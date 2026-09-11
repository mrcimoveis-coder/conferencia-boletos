import io
import re
import unicodedata
import openpyxl
import pypdf
import streamlit as st

def normalizar_nome(texto):
    """Remove acentos, pontuação, caracteres especiais e converte para maiúsculas."""
    if not texto:
        return ""
    texto = unicodedata.normalize("NFD", str(texto))
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    texto = texto.upper()
    texto = re.sub(r"[^A-Z0-9\s]", " ", texto)
    return re.sub(r"\s+", " ", texto).strip()

def normalizar_valor(valor_raw):
    """Converte o valor monetário removendo R$, separador de milhar e formatando para float com 2 casas decimais."""
    if valor_raw is None:
        return None
    if isinstance(valor_raw, (int, float)):
        return round(float(valor_raw), 2)
    s = str(valor_raw).replace("R$", "").replace(".", "").replace(",", ".").strip()
    try:
        return round(float(s), 2)
    except ValueError:
        return None

# Interface Web Streamlit
st.set_page_config(page_title="Conferência de Boletos", page_icon="📑", layout="centered")

# Exibição da Logo da MRC (busca primeiro a imagem do repositório da intranet)
try:
    st.image("https://raw.githubusercontent.com/mrcimoveis-coder/intranet/main/logo.jpeg", width=180)
except Exception:
    try:
        st.image("logo.jpeg", width=180)
    except Exception:
        pass

st.title("📑 Conferência Final de Boletos")
st.write("Validação automática do PDF da Crítica de Boletos contra a aba CONF do Excel.")

file_excel = st.file_uploader("1. Selecione a Planilha Excel (Aba CONF)", type=["xlsx"])
file_pdf = st.file_uploader("2. Selecione o PDF (Crítica de Boletos)", type=["pdf"])

if file_excel and file_pdf:
    if st.button("🚀 Processar Conferência", type="primary"):
        with st.spinner("Processando e aplicando regras de validação..."):
            # 1. Leitura e extração do campo 'Total da cobrança' no PDF
            reader = pypdf.PdfReader(file_pdf)
            texto_completo = "\n".join(page.extract_text() or "" for page in reader.pages)

            pdf_dados = {}
            blocos = texto_completo.split("Locatário:")
            for bloco in blocos[1:]:
                linhas = bloco.strip().split("\n")
                nome_raw = re.sub(r"\s*\(\d+\)\s*$", "", linhas[0].strip())
                nome_norm = normalizar_nome(nome_raw)
                
                match_total = re.search(r"Total\s+da\s+cobrança:\s*([\d\.\s]+,\d{2})", bloco)
                if match_total:
                    val = normalizar_valor(match_total.group(1))
                    if val is not None:
                        if nome_norm not in pdf_dados:
                            pdf_dados[nome_norm] = {'nome_original': nome_raw, 'valores': []}
                        pdf_dados[nome_norm]['valores'].append(val)

            # 2. Leitura do Excel (Aba CONF, linha 4 até o final)
            wb = openpyxl.load_workbook(file_excel)
            aba_alvo = next((name for name in wb.sheetnames if name.strip().upper() == "CONF"), wb.sheetnames[0])
            ws = wb[aba_alvo]

            locatarios_validados_excel = set()

            for row in range(4, ws.max_row + 1):
                locatario_val = ws.cell(row=row, column=10).value  # Coluna J: Nome do locatário
                if not locatario_val or str(locatario_val).strip() == "":
                    continue

                val_esperado = normalizar_valor(ws.cell(row=row, column=3).value)  # Coluna C: Valor esperado
                locatario_norm = normalizar_nome(str(locatario_val))
                locatarios_validados_excel.add(locatario_norm)

                resultado = "ERRO"
                if locatario_norm in pdf_dados and val_esperado is not None:
                    valores_pdf = pdf_dados[locatario_norm]['valores']
                    # Tratamento de duplicidade: deve haver exatamente 1 correspondência monetária exata
                    matches = [v for v in valores_pdf if abs(v - val_esperado) < 1e-4]
                    if len(matches) == 1:
                        resultado = "OK"

                ws.cell(row=row, column=9, value=resultado)  # Coluna I: Resultado (OK ou ERRO)

            # 3. Aba adicional: Sobra PDF
            if "Sobra PDF" in wb.sheetnames:
                del wb["Sobra PDF"]
            
            ws_sobra = wb.create_sheet("Sobra PDF")
            ws_sobra.append(["Locatário no PDF", "Valor Encontrado no PDF"])
            
            for norm_pdf, dados_pdf in pdf_dados.items():
                if norm_pdf not in locatarios_validados_excel:
                    for v in dados_pdf['valores']:
                        ws_sobra.append([dados_pdf['nome_original'], v])

            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            
            nome_saida = f"{file_excel.name.replace('.xlsx', '')}_CONFERIDO.xlsx"

        st.success("✅ Conferência concluída com sucesso!")
        
        st.download_button(
            label="📥 Baixar Planilha Conferida",
            data=buffer,
            file_name=nome_saida,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
